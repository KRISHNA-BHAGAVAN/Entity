import json
import hashlib
import time
from typing import List, Dict, Any, TypedDict, Optional, Tuple
from functools import lru_cache
from dotenv import load_dotenv
from langchain.chat_models import init_chat_model
from langchain_core.messages import SystemMessage, HumanMessage
from langgraph.graph import StateGraph, START, END
import redis
from rapidfuzz import fuzz
import tiktoken
import re
import os


# --------------------------------------------------------------------------
# MARKDOWN LOCATION TRACKING (Primary Method)
# --------------------------------------------------------------------------
def find_reference_locations_in_markdown(
    markdown_content: str, reference: str, filename: str
) -> List[Dict[str, Any]]:
    """
    Find ALL exact locations of `reference` in MARKDOWN with precise positions.
    Perfect for frontend highlighting in markdown previews.
    """
    if not markdown_content or not reference:
        return []

    locations = []
    lines = markdown_content.split("\n")

    for line_idx, line in enumerate(lines):
        start_pos = 0
        while True:
            pos = line.find(reference, start_pos)
            if pos == -1:
                break

            # GLOBAL positions (full markdown)
            global_char_start = sum(len(l) + 1 for l in lines[:line_idx]) + pos
            global_char_end = global_char_start + len(reference)

            # LINE positions (within this line only)
            line_char_start = pos
            line_char_end = pos + len(reference)

            # Determine context type
            line_type = "paragraph"
            line_stripped = line.strip()
            if line_stripped.startswith("|"):  # Markdown table
                line_type = "table_row"
            elif re.match(r"^#{1,6}\s", line):  # Header
                line_type = "header"
            elif line_stripped.startswith("- ") or line_stripped.startswith("* "):  # List
                line_type = "list_item"
            elif line_stripped.startswith(">"):  # Blockquote
                line_type = "blockquote"

            locations.append(
                {
                    "filename": filename,
                    "type": line_type,
                    "line_index": line_idx,
                    "char_start": global_char_start,
                    "char_end": global_char_end,
                    "line_char_start": line_char_start,
                    "line_char_end": line_char_end,
                    "text": reference,
                    "context_line": line_stripped[:100],  # Truncated for readability
                }
            )
            start_pos = pos + 1  # Allow overlapping matches

    return locations


def count_reference_in_markdown(markdown_content: str, reference: str) -> int:
    """Count occurrences in markdown (for backward compatibility)."""
    locations = find_reference_locations_in_markdown(markdown_content, reference, "")
    return len(locations)


# --------------------------------------------------------------------------
# DOCX FALLBACK (Optional)
# --------------------------------------------------------------------------


def normalize_text(text: str) -> str:
    if not text:
        return text
    return (
        text.replace("\u201c", '"')
        .replace("\u201d", '"')
        .replace("\u2018", "'")
        .replace("\u2019", "'")
        .replace("\u2013", "-")
        .replace("\u2014", "-")
        .replace("\u00a0", " ")
        .replace("\u2026", ".")
    )


def find_reference_locations_in_docx(
    doc_path: str, reference: str
) -> List[Dict[str, Any]]:
    """Fallback for DOCX files - only used if doc_paths provided."""
    if not os.path.exists(doc_path):
        return []

    try:
        from docx import Document as DocxDocument

        doc = DocxDocument(doc_path)
        filename = os.path.basename(doc_path)
        locations = []

        # Paragraphs only (fast)
        for para_idx, para in enumerate(doc.paragraphs):
            text = para.text or ""
            pos = text.find(reference)
            if pos != -1:
                locations.append(
                    {
                        "filename": filename,
                        "type": "paragraph",
                        "paragraph_index": para_idx,
                        "char_start": pos,
                        "char_end": pos + len(reference),
                        "text": reference,
                    }
                )
        return locations
    except Exception:
        return []


# ------------------------------------------------------------------------------
# TOKEN COUNTING
# ------------------------------------------------------------------------------


_ENCODING = tiktoken.get_encoding("cl100k_base")


def get_token_count(text: str) -> int:
    try:
        return len(_ENCODING.encode(text))
    except Exception:
        return len(text) // 4 + 1


# ------------------------------------------------------------------------------
# ENV + LLM + REDIS CACHE
# ------------------------------------------------------------------------------


load_dotenv(override=True)

# LLM will be initialized per-request using BYOK
llm = None

try:
    redis_client = redis.Redis(
        host="localhost", port=6379, db=0, decode_responses=True
    )
    redis_client.ping()
    USE_REDIS = True
    print("✅ Redis cache enabled")
except Exception:
    redis_client = None
    USE_REDIS = False
    print("⚠️ Using in-memory cache")


# ------------------------------------------------------------------------------
# STATE
# ------------------------------------------------------------------------------


class SchemaDiscoveryState(TypedDict, total=False):
    documents: List[Tuple[str, str]]  # (filename, markdown_content)
    doc_paths: List[str]  # Optional DOCX paths
    cache_key: Optional[str]
    partial_schemas: List[Dict[str, Any]]
    final_schema: Optional[Dict[str, Any]]
    stats: Dict[str, Any]
    # user-provided extraction instructions (optional)
    user_instructions: Optional[str]
    # BYOK integration
    user_id: Optional[str]
    jwt_token: Optional[str]
    llm_instance: Optional[Any]


INITIAL_STATS: Dict[str, Any] = {
    "cache_hit": False,
    "processing_time": 0.0,
    "docs_processed": 0,
    "total_fields": 0,
    "total_locations": 0,
    "sections_created": 0,
    "total_chars_processed": 0,
    "merge_time": 0.0,
    "llm": {
        "calls": [],
        "summary": {
            "llm_calls": 0,
            "total_input_tokens": 0,
            "total_output_tokens": 0,
            "total_tokens": 0,
            "avg_input_tokens_per_call": 0,
            "avg_output_tokens_per_call": 0,
        },
    },
}


# ------------------------------------------------------------------------------
# PROMPTS
# ------------------------------------------------------------------------------


SCHEMA_DISCOVERY_PROMPT = """
You extract editable fields from documents.

Document: {filename}

User instructions (may be empty):
{user_instructions}

Task:
- If user instructions are provided: extract ONLY fields that match or are clearly implied by them.
- If no instructions are provided: extract fields users are likely to change over time.
- Ignore all tabular data.
- Exclude fields with no references.

Field rules:
- Top-level keys must be stable snake_case.
- Each field contains:
  - label: human-readable name
  - references: all exact text spans referring to the same fact

Reference rules:
- References must be exact substrings from the document.
- Group all textual variants of the same fact into one field.
- Deduplicate references within a field.
- Do not repeat the same text across different fields.
- Do not include input placeholders like "________".
Output:
- STRICT JSON only
- One flat JSON object
- Keys = field names
- Values = field objects
- No markdown, prose, or comments
"""

CONSOLIDATION_PROMPT = """
You are consolidating extracted document fields.

Input JSON:
{fields_json}

Task:
- Identify fields that refer to the same real-world entity or fact,
  even if they appear under different roles, names, formats, or contexts.
- Merge such fields into a single canonical field.
- Choose a stable snake_case key that best represents the entity or fact.
- Preserve ALL original references under the merged field.
- Do not invent new information.

Rules:
- If two or more keys clearly refer to the same person, event, organization, venue, date range, ID, or other real-world fact,
  merge them into one field.
- When merging:
    - Pick a clear, generic snake_case key (e.g. primary_applicant_name, main_event_dates, organization_name).
    - Combine all references from all merged fields (deduplicated exact strings).
- If multiple fields represent different granular views of the same underlying concept (e.g. range vs components), prefer a single higher-level field unless information would be lost
- Minor spelling differences, abbreviated names, titles, or role-based mentions may still refer to the same entity

Output:
- STRICT JSON only.
- Flat object.
- Keys = canonical field names (snake_case).
- Values are objects with:
    - label      : human-readable label
    - references : array of exact original strings (from the input)
- No markdown, no prose, no comments.
"""


# ------------------------------------------------------------------------------
# UTILS
# ------------------------------------------------------------------------------


@lru_cache(maxsize=128)
def normalize_key(key: str) -> str:
    return key.lower().replace(" ", "_").replace("-", "_")


def fuzzy_dedupe_references(
    references: List[str], threshold: int = 85
) -> List[str]:
    if len(references) <= 1:
        return references
    deduped = []
    for ref in sorted(references, key=len, reverse=True):
        is_duplicate = any(
            fuzz.ratio(ref, existing) >= threshold for existing in deduped
        )
        if not is_duplicate:
            deduped.append(ref)
    return deduped


CACHE: Dict[str, Any] = {}


def get_cache(key: str) -> Optional[Dict[str, Any]]:
    if USE_REDIS:
        try:
            cached = redis_client.get(f"schema:{key}")
            return json.loads(cached) if cached else None
        except Exception:
            return None
    return CACHE.get(key)


def set_cache(key: str, value: Dict[str, Any], ttl: int = 3600) -> None:
    if USE_REDIS:
        try:
            redis_client.setex(f"schema:{key}", ttl, json.dumps(value))
            return
        except Exception:
            pass
    CACHE[key] = value


def track_llm_usage(
    stats: Dict[str, Any],
    input_tokens: int,
    output_tokens: int,
    prompt_chars: int,
) -> Dict[str, Any]:
    llm_stats = stats.setdefault(
        "llm",
        {
            "calls": [],
            "summary": {
                "llm_calls": 0,
                "total_input_tokens": 0,
                "total_output_tokens": 0,
                "total_tokens": 0,
                "avg_input_tokens_per_call": 0,
                "avg_output_tokens_per_call": 0,
            },
        },
    )

    call_stats = {
        "timestamp": time.time(),
        "input_tokens": input_tokens,
        "output_tokens": output_tokens,
        "prompt_chars": prompt_chars,
        "total_tokens": input_tokens + output_tokens,
    }
    llm_stats["calls"].append(call_stats)

    summary = llm_stats["summary"]
    summary["llm_calls"] += 1
    summary["total_input_tokens"] += input_tokens
    summary["total_output_tokens"] += output_tokens
    summary["total_tokens"] = (
        summary["total_input_tokens"] + summary["total_output_tokens"]
    )
    if summary["llm_calls"] > 0:
        summary["avg_input_tokens_per_call"] = (
            summary["total_input_tokens"] // summary["llm_calls"]
        )
        summary["avg_output_tokens_per_call"] = (
            summary["total_output_tokens"] // summary["llm_calls"]
        )
    return stats


# ------------------------------------------------------------------------------
# NODES
# ------------------------------------------------------------------------------


def cache_check(state: SchemaDiscoveryState) -> Dict[str, Any]:
    stats = state.get("stats", INITIAL_STATS.copy())

    parts = []
    total_chars = 0
    for filename, markdown in state["documents"]:
        parts.append(f"{filename}:{markdown[:1000]}")
        total_chars += len(markdown)
    for p in state.get("doc_paths", []):
        parts.append(f"docx:{p}")

    # include user instructions (or empty) in cache key so different instructions
    # over same docs don't collide.
    user_instr = (state.get("user_instructions") or "").strip()
    parts.append(f"user_instructions:{user_instr[:500]}")

    content_hash = hashlib.sha256("".join(parts).encode()).hexdigest()
    stats["total_chars_processed"] = total_chars

    cached = get_cache(content_hash)
    if cached:
        print(f"✅ CACHE HIT: {content_hash[:8]}")
        stats["cache_hit"] = True
        stats["processing_time"] = 0.001
        stats["total_locations"] = cached.get("stats", {}).get(
            "total_locations", 0
        )
        return {
            "final_schema": cached["schema"],
            "stats": stats,
        }

    print(f"🔄 CACHE MISS: {content_hash[:8]}")
    return {"cache_key": content_hash, "stats": stats}


def map_discover_schema(state: SchemaDiscoveryState) -> Dict[str, Any]:
    """
    Phase A — Raw Discovery.
    Over-extract, do not globally dedupe, allow overlaps.
    """
    partials: List[Dict[str, Any]] = []
    stats = state.get("stats", INITIAL_STATS.copy())
    
    # Get LLM instance from state (BYOK)
    llm_instance = state.get("llm_instance")
    if not llm_instance:
        raise ValueError("No LLM instance available")

    user_instructions_raw = state.get("user_instructions") or ""
    user_instructions_for_prompt = user_instructions_raw.strip()

    for i, (filename, md_raw) in enumerate(state["documents"]):
        raw_length = len(md_raw)
        print(f"\n🔍 DOC {i+1} ({filename}): RAW={raw_length} chars")

        if raw_length < 50:
            print("  ⏭️ SKIP: too short")
            continue

        try:
            content = md_raw[:8000]
            print(f"  📤 Sending {len(content)} chars to LLM...")

            prompt = SCHEMA_DISCOVERY_PROMPT.format(
                filename=filename,
                user_instructions=user_instructions_for_prompt,
            )
            full_prompt = prompt + "\n\n" + content

            prompt_tokens = get_token_count(full_prompt)
            start_time = time.time()

            response = llm_instance.invoke(
                [
                    SystemMessage(content=prompt),
                    HumanMessage(content=content),
                ]
            )

            response_time = time.time() - start_time
            output_tokens = get_token_count(response.content)

            print(
                f"  📥 LLM RAW RESPONSE ({len(response.content)} chars, ~{output_tokens} tokens)"
            )
            print(f"     {repr(response.content[:200])}...")

            response_text = response.content.strip()
            start = response_text.find("{")
            end = response_text.rfind("}") + 1
            json_str = response_text[start:end] if start != -1 else None

            if not json_str:
                print("  ❌ NO JSON FOUND!")
                continue

            partial_schema = json.loads(json_str)
            print(
                f"  ✅ PARSED: {len(partial_schema)} keys: {list(partial_schema.keys())}"
            )

            # Per-field internal dedupe only (no global dedupe, allow overlaps)
            for field_key, field_val in partial_schema.items():
                if isinstance(field_val, dict) and "references" in field_val:
                    refs = field_val.get("references") or []
                    # exact and fuzzy dedupe within the field
                    unique_refs = list(dict.fromkeys(refs))
                    unique_refs = fuzzy_dedupe_references(unique_refs)
                    field_val["references"] = unique_refs
                    field_val["source_filename"] = filename

            partials.append(partial_schema)
            print(f"  🎉 ADDED DOC {i+1} ({filename}) - {response_time:.2f}s")

            stats = track_llm_usage(
                stats,
                input_tokens=prompt_tokens,
                output_tokens=output_tokens,
                prompt_chars=len(full_prompt),
            )

        except Exception as e:
            print(f"  💥 ERROR DOC {i+1}: {e}")
            continue

    stats["docs_processed"] = len(partials)
    print(f"\n🎯 TOTAL PARTIALS: {len(partials)}")
    return {"partial_schemas": partials, "stats": stats}


def merge_schemas_enhanced(state: SchemaDiscoveryState) -> Dict[str, Any]:
    """
    Phase A continued — structural merge only.
    No semantic consolidation; just union of fields + references across docs.
    """
    start_time = time.time()
    merged: Dict[str, Any] = {}

    for schema in state["partial_schemas"]:
        print(f"🔍 Merging schema: {len(schema)} keys")

        if "$schema" in schema:
            schema = schema.get("properties", {})

        section_key = "document_fields"
        section_val = {
            "label": "Document Fields",
            "fields": schema,
        }

        merged_section = merged.setdefault(
            section_key,
            {
                "label": "Document Fields",
                "fields": {},
                "doc_frequency": 0,
            },
        )
        merged_section["doc_frequency"] += 1

        for field_key, field_val in section_val["fields"].items():
            if not isinstance(field_val, dict):
                continue

            merged_field = merged_section["fields"].setdefault(
                field_key, {"label": field_val.get("label"), "references": []}
            )

            # union of references (exact-dedup only here)
            existing_refs = merged_field.get("references") or []
            new_refs = field_val.get("references") or []
            merged_refs = list(dict.fromkeys(existing_refs + new_refs))
            merged_field["references"] = merged_refs

            merged_field["doc_frequency"] = (
                merged_field.get("doc_frequency", 0) + 1
            )

            source_files = merged_field.setdefault("source_files", [])
            source_file = field_val.get("source_filename")
            if source_file and source_file not in source_files:
                source_files.append(source_file)

    # Sort fields by importance: doc_frequency then number of references
    for section in merged.values():
        fields = section.get("fields", {})
        if isinstance(fields, dict):
            section["fields"] = dict(
                sorted(
                    fields.items(),
                    key=lambda x: (
                        x[1].get("doc_frequency", 0),
                        len(x[1].get("references", [])),
                    ),
                    reverse=True,
                )
            )

    processing_time = time.time() - start_time
    total_fields = sum(len(s.get("fields", {})) for s in merged.values())

    stats = state["stats"].copy()
    stats.update(
        {
            "processing_time": processing_time,
            "total_fields": total_fields,
            "sections_created": len(merged),
            "merge_time": processing_time,
        }
    )

    print(
        f"🎉 FINAL MERGE (pre-consolidation): {total_fields} fields from {len(state['partial_schemas'])} docs"
    )
    return {
        "final_schema": merged,
        "stats": stats,
    }


def consolidate_entities_llm(state: SchemaDiscoveryState) -> Dict[str, Any]:
    """
    Phase B — Entity/Fact Consolidation.
    Take merged.fields -> compact text-only structure -> LLM -> unified canonical fields.
    """
    final_schema = state.get("final_schema", {})
    if not final_schema:
        return {}
    
    # Get LLM instance from state (BYOK)
    llm_instance = state.get("llm_instance")
    if not llm_instance:
        raise ValueError("No LLM instance available")

    document_fields_section = final_schema.get("document_fields", {})
    fields = document_fields_section.get("fields", {}) or {}

    # Build compact structure: { field_key: [references...] }
    compact_fields: Dict[str, List[str]] = {}
    for field_key, field_val in fields.items():
        refs = field_val.get("references") or []
        if not refs:
            continue
        compact_fields[field_key] = refs

    compact_payload = {"fields": compact_fields}
    fields_json = json.dumps(compact_payload, ensure_ascii=False)

    prompt = CONSOLIDATION_PROMPT.format(fields_json=fields_json)
    prompt_tokens = get_token_count(prompt)

    start_time = time.time()
    response = llm_instance.invoke(
        [
            SystemMessage(content=prompt),
            HumanMessage(content="Return only the consolidated JSON."),
        ]
    )
    response_time = time.time() - start_time
    output_tokens = get_token_count(response.content)

    print(
        f"🔁 CONSOLIDATION LLM RESPONSE ({len(response.content)} chars, ~{output_tokens} tokens, {response_time:.2f}s)"
    )
    print(f"    {repr(response.content[:200])}...")

    response_text = response.content.strip()
    start = response_text.find("{")
    end = response_text.rfind("}") + 1
    json_str = response_text[start:end] if start != -1 else None

    if not json_str:
        print("  ❌ NO JSON FOUND DURING CONSOLIDATION — keeping original merged schema")
        consolidated_fields = fields
    else:
        consolidated_fields = json.loads(json_str)

    # Re-wrap into the same "document_fields" section structure, preserving source_files if possible
    new_fields: Dict[str, Any] = {}
    for canon_key, canon_val in consolidated_fields.items():
        # try to reuse some label, otherwise fallback to provided label
        label = canon_val.get("label") or canon_key.replace("_", " ").title()
        refs = canon_val.get("references") or []

        # Collect source_files from original fields where the references came from
        source_files: List[str] = []
        for original_key, original_field in fields.items():
            orig_refs = original_field.get("references") or []
            if any(r in orig_refs for r in refs):
                for sf in original_field.get("source_files", []):
                    if sf not in source_files:
                        source_files.append(sf)

        new_fields[canon_key] = {
            "label": label,
            "references": refs,
            "source_files": source_files,
            # doc_frequency is approximate but useful: how many original fields contributed
            "doc_frequency": sum(
                1
                for original_key, original_field in fields.items()
                if any(
                    r in (original_field.get("references") or [])
                    for r in refs
                )
            ),
        }

    final_schema["document_fields"]["fields"] = new_fields

    stats = state["stats"].copy()
    stats = track_llm_usage(
        stats,
        input_tokens=prompt_tokens,
        output_tokens=output_tokens,
        prompt_chars=len(prompt),
    )

    print(
        f"✅ CONSOLIDATION DONE: {len(new_fields)} canonical fields from {len(fields)} raw fields"
    )
    return {"final_schema": final_schema, "stats": stats}


def compute_frequencies_and_locations(state: SchemaDiscoveryState) -> Dict[str, Any]:
    """
    Phase C — Re-attach locations.
    Compute frequencies AND precise markdown locations for frontend highlighting.
    """
    final_schema = state.get("final_schema", {})
    documents = state.get("documents", [])
    doc_paths = state.get("doc_paths", [])

    if not final_schema or not documents:
        return {}

    path_by_basename = {os.path.basename(p): p for p in doc_paths}
    document_fields = final_schema.get("document_fields", {})
    fields = document_fields.get("fields", {})

    print("📍 COMPUTING MARKDOWN LOCATIONS + FREQUENCIES:")
    total_locations = 0

    for field_key, field_val in fields.items():
        refs = field_val.get("references", [])
        source_files = field_val.get("source_files", [])
        all_locations = []
        total_freq = 0

        for ref in refs:
            if not ref:
                continue

            # PRIMARY: Search in markdown documents
            for filename, markdown_content in documents:
                if not source_files or any(
                    fname in filename for fname in source_files
                ):
                    locations = find_reference_locations_in_markdown(
                        markdown_content, ref, filename
                    )
                    all_locations.extend(locations)
                    total_freq += len(locations)
                    if locations:
                        print(
                            f"  📍 {field_key}: '{ref}' → {len(locations)} locs in {filename}"
                        )

            # FALLBACK: DOCX files (if provided)
            for fname in source_files:
                full_path = path_by_basename.get(fname) or next(
                    (p for p in doc_paths if p.endswith(fname)), None
                )
                if full_path:
                    docx_locations = find_reference_locations_in_docx(
                        full_path, ref
                    )
                    all_locations.extend(docx_locations)
                    total_freq += len(docx_locations)

        # Store COMPLETE locations for frontend
        field_val["locations"] = all_locations
        field_val["frequency"] = total_freq
        field_val["location_count"] = len(all_locations)
        total_locations += len(all_locations)

    stats = state["stats"].copy()
    stats["total_locations"] = total_locations
    print(f"✅ {total_locations} TOTAL LOCATIONS computed across all fields")
    return {"final_schema": final_schema, "stats": stats}


def cache_store(state: SchemaDiscoveryState) -> Dict[str, Any]:
    if state.get("cache_key"):
        cache_data = {
            "schema": state["final_schema"],
            "stats": state.get("stats", {}),
        }
        set_cache(state["cache_key"], cache_data)
        print(
            f"💾 CACHED: {state['cache_key'][:8]} with {state['stats'].get('total_locations', 0)} locations"
        )
    return {}


# ------------------------------------------------------------------------------
# LANGGRAPH
# ------------------------------------------------------------------------------


graph = StateGraph(SchemaDiscoveryState)

graph.add_node("cache_check", cache_check)
graph.add_node("map_discover_schema", map_discover_schema)
graph.add_node("merge_schemas", merge_schemas_enhanced)
graph.add_node("consolidate_entities", consolidate_entities_llm)
graph.add_node(
    "compute_frequencies_and_locations", compute_frequencies_and_locations
)
graph.add_node("cache_store", cache_store)

graph.add_edge(START, "cache_check")
graph.add_conditional_edges(
    "cache_check", lambda s: END if s.get("final_schema") else "map_discover_schema"
)
graph.add_edge("map_discover_schema", "merge_schemas")
graph.add_edge("merge_schemas", "consolidate_entities")
graph.add_edge("consolidate_entities", "compute_frequencies_and_locations")
graph.add_edge("compute_frequencies_and_locations", "cache_store")
graph.add_edge("cache_store", END)

schema_discovery_workflow = graph.compile()
